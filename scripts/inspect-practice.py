from pathlib import Path
import zipfile,json,subprocess,hashlib,io,re
import pymupdf
from docx import Document
import openpyxl
root=Path(__file__).resolve().parents[1]
inventory=json.loads((root/'reports/sources/inventory.json').read_text(encoding='utf-8'))
out=root/'reports/practice-inputs-v3';out.mkdir(parents=True,exist_ok=True)
records=[]
for r in inventory:
 if 'entries' not in r or not any(f'09.02.07-{n}-' in r['name'] for n in (2,3,5)):continue
 slug={2:'programmer',3:'web',5:'is'}[next(n for n in (2,3,5) if f'09.02.07-{n}-' in r['name'])]
 dest=out/slug/str(r['year']);dest.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(r['path']) as z:
  for e in z.infolist():
   if not e.filename.startswith('БУ/') or e.is_dir():continue
   assert '..' not in Path(e.filename).parts
   target=dest/e.filename;target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(z.read(e))
 for archive in list(dest.rglob('*.rar'))+list(dest.rglob('*.zip')):
  unpack=archive.parent/(archive.stem+'_files');unpack.mkdir(exist_ok=True)
  if archive.suffix=='.rar':
   subprocess.run(['C:/Program Files/WinRAR/UnRAR.exe','x','-o+','-idq',str(archive),str(unpack)+'\\'],check=True,creationflags=subprocess.CREATE_NO_WINDOW)
  else:
   with zipfile.ZipFile(archive) as z:
    for e in z.infolist():
     name=e.filename
     if not e.flag_bits & 0x800:
      try:name=name.encode('cp437').decode('cp866')
      except UnicodeError:pass
     assert '..' not in Path(name).parts and not Path(name).is_absolute()
     target=unpack/name
     if e.is_dir():target.mkdir(parents=True,exist_ok=True)
     else:target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(z.read(e))
 for f in dest.rglob('*'):
  if not f.is_file():continue
  rec={'path':str(f.relative_to(out)),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()}
  if f.suffix.lower()=='.pdf':
   d=pymupdf.open(f);rec['text']='\n'.join(p.get_text() for p in d)
  elif f.suffix.lower()=='.docx':
   d=Document(f);rec['text']='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows)
  elif f.suffix.lower()=='.xlsx':
   w=openpyxl.load_workbook(f,data_only=False,read_only=True);rec['sheets']=[{'name':s.title,'rows':s.max_row,'cols':s.max_column,'sample':[[str(v) if v is not None else '' for v in row] for row in s.iter_rows(values_only=True)]} for s in w];w.close()
  elif f.suffix.lower() in ('.json','.txt','.sql'):
   rec['text']=f.read_text(encoding='utf-8-sig',errors='replace')
  records.append(rec)
(out/'inspection.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print('Inspected',len(records),'files')
for r in records:
 if r['path'].startswith('is/') and ('text' in r or 'sheets' in r):print(r['path'],len(r.get('text','')),[(s['name'],s['rows'],s['cols']) for s in r.get('sheets',[])])
